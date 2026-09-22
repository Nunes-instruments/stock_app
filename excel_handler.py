from pathlib import Path
from datetime import datetime
from io import BytesIO
import hashlib
import shutil

import pandas as pd
from openpyxl import Workbook, load_workbook

from database import (
    get_connection,
    get_current_timestamp,
    get_active_branch_key,
    get_branch_slug,
)
from stock_service import add_stock, get_inventory_products
from runtime_paths import UPLOAD_DIR, PROCESSED_DIR, EXPORT_DIR


# =========================================================
# PATHS
# =========================================================

BASE_DIR = Path(__file__).resolve().parent

MASTER_STOCK_FILE = EXPORT_DIR / "Master_Stock.xlsx"


def get_master_stock_file():
    """Return a separate master Excel file for the active branch."""
    branch_key = get_active_branch_key()
    if branch_key == "main":
        return MASTER_STOCK_FILE
    return EXPORT_DIR / f"Master_Stock_{get_branch_slug(branch_key)}.xlsx"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
EXPORT_DIR.mkdir(parents=True, exist_ok=True)


# =========================================================
# EXCEL COLUMN SETTINGS
# =========================================================

REQUIRED_COLUMNS = [
    "Product ID",
    "Category",
    "Product Name",
    "Quantity",
]

OPTIONAL_COLUMNS = [
    "Brand",
    "Model",
    "Unit",
    "Location",
    "Stock Date",
    "Remarks",
]


# =========================================================
# GENERAL CLEANING
# =========================================================

def normalize_column_name(value):
    if value is None:
        return ""

    return str(value).strip()


def clean_value(value):
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    return str(value).strip()


def clean_quantity(value):
    if value is None:
        raise ValueError("Quantity is missing.")

    try:
        if pd.isna(value):
            raise ValueError("Quantity is missing.")
    except TypeError:
        pass

    if str(value).strip() == "":
        raise ValueError("Quantity is missing.")

    try:
        quantity = float(value)
    except (TypeError, ValueError):
        raise ValueError("Quantity must be a valid number.")

    if quantity <= 0:
        raise ValueError("Quantity must be greater than zero.")

    return quantity


# =========================================================
# STOCK DATE HANDLING
# =========================================================

def clean_stock_date(value):
    """
    Convert Excel/Pandas dates into YYYY-MM-DD.
    If blank, today's date is used.
    """

    if value is None:
        return datetime.now().strftime("%Y-%m-%d")

    try:
        if pd.isna(value):
            return datetime.now().strftime("%Y-%m-%d")
    except Exception:
        pass

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    text_value = str(value).strip()

    if not text_value:
        return datetime.now().strftime("%Y-%m-%d")

    try:
        parsed_date = pd.to_datetime(text_value)

        return parsed_date.strftime("%Y-%m-%d")

    except Exception:
        return text_value


# =========================================================
# FILE HASH / DUPLICATE FILE PROTECTION
# =========================================================

def generate_file_hash(file_path):
    sha256 = hashlib.sha256()

    with open(file_path, "rb") as file:
        while True:
            chunk = file.read(1024 * 1024)

            if not chunk:
                break

            sha256.update(chunk)

    return sha256.hexdigest()


def ensure_import_hash_column():
    """
    Older versions of the database may not have file_hash.
    Add it automatically if missing.
    """

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        """
        PRAGMA table_info(import_history)
        """
    )

    columns = [
        row["name"]
        for row in cursor.fetchall()
    ]

    if "file_hash" not in columns:
        cursor.execute(
            """
            ALTER TABLE import_history
            ADD COLUMN file_hash TEXT
            """
        )

    conn.commit()
    conn.close()


def is_file_already_imported(file_hash):
    ensure_import_hash_column()

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        """
        SELECT *
        FROM import_history
        WHERE file_hash = ?
          AND status IN (
              'Completed',
              'Completed With Errors'
          )
        ORDER BY id DESC
        LIMIT 1
        """,
        (file_hash,),
    )

    row = cursor.fetchone()
    conn.close()

    return dict(row) if row else None


# =========================================================
# SAVE STREAMLIT UPLOAD
# =========================================================

def save_uploaded_file(uploaded_file):
    """
    Save a Streamlit uploaded file into uploads/.
    """

    safe_name = Path(uploaded_file.name).name

    destination = UPLOAD_DIR / safe_name

    with open(destination, "wb") as file:
        file.write(uploaded_file.getbuffer())

    return destination


# =========================================================
# READ EXCEL
# =========================================================

def read_excel_file(file_path):
    file_path = Path(file_path)

    if not file_path.exists():
        raise ValueError("Excel file was not found.")

    if file_path.suffix.lower() != ".xlsx":
        raise ValueError(
            "Please use .xlsx Excel files."
        )

    try:
        dataframe = pd.read_excel(
            file_path,
            dtype=object,
            engine="openpyxl",
        )

    except Exception as error:
        raise ValueError(
            f"Unable to read Excel file: {error}"
        )

    dataframe.columns = [
        normalize_column_name(column)
        for column in dataframe.columns
    ]

    dataframe = dataframe.dropna(
        how="all"
    ).reset_index(drop=True)

    return dataframe


# =========================================================
# COLUMN VALIDATION
# =========================================================

def validate_excel_columns(dataframe):
    missing_columns = []

    for required_column in REQUIRED_COLUMNS:

        if required_column not in dataframe.columns:
            missing_columns.append(
                required_column
            )

    if missing_columns:

        missing_text = ", ".join(
            missing_columns
        )

        raise ValueError(
            "Missing required Excel column(s): "
            f"{missing_text}"
        )

    return True


# =========================================================
# EXCEL PREVIEW
# =========================================================

def prepare_preview(dataframe):
    validate_excel_columns(
        dataframe
    )

    preview_rows = []

    for dataframe_index, row in dataframe.iterrows():

        excel_row = dataframe_index + 2

        product_id = clean_value(
            row.get("Product ID")
        )

        category = clean_value(
            row.get("Category")
        )

        product_name = clean_value(
            row.get("Product Name")
        )

        errors = []

        if not product_id:
            errors.append(
                "Product ID missing"
            )

        if not category:
            errors.append(
                "Category missing"
            )

        if not product_name:
            errors.append(
                "Product Name missing"
            )

        try:
            quantity = clean_quantity(
                row.get("Quantity")
            )

        except ValueError as error:
            quantity = ""
            errors.append(
                str(error)
            )

        preview_rows.append(
            {
                "Excel Row": excel_row,
                "Product ID": product_id,
                "Category": category,
                "Product Name": product_name,
                "Brand": clean_value(
                    row.get("Brand")
                ),
                "Model": clean_value(
                    row.get("Model")
                ),
                "Quantity": quantity,
                "Unit": (
                    clean_value(
                        row.get("Unit")
                    )
                    or "Nos"
                ),
                "Location": clean_value(
                    row.get("Location")
                ),
                "Stock Date": clean_stock_date(
                    row.get("Stock Date")
                ),
                "Remarks": clean_value(
                    row.get("Remarks")
                ),
                "Validation Status": (
                    "Ready"
                    if not errors
                    else "Failed"
                ),
                "Validation Error": "; ".join(
                    errors
                ),
            }
        )

    return pd.DataFrame(
        preview_rows
    )


# =========================================================
# IMPORT HISTORY
# =========================================================

def generate_import_number(import_id):
    return f"IMP{import_id:06d}"


def create_import_history(
    file_name,
    original_file_name,
    total_rows,
    file_hash,
):
    ensure_import_hash_column()

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        """
        INSERT INTO import_history (
            import_number,
            file_name,
            original_file_name,
            total_rows,
            success_rows,
            failed_rows,
            total_quantity_added,
            status,
            imported_at,
            file_hash
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            None,
            file_name,
            original_file_name,
            total_rows,
            0,
            0,
            0,
            "Processing",
            get_current_timestamp(),
            file_hash,
        ),
    )

    import_id = cursor.lastrowid

    import_number = generate_import_number(
        import_id
    )

    cursor.execute(
        """
        UPDATE import_history
        SET import_number = ?
        WHERE id = ?
        """,
        (
            import_number,
            import_id,
        ),
    )

    conn.commit()
    conn.close()

    return import_id, import_number


def update_import_history(
    import_id,
    success_rows,
    failed_rows,
    total_quantity_added,
    status,
):
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        """
        UPDATE import_history
        SET success_rows = ?,
            failed_rows = ?,
            total_quantity_added = ?,
            status = ?
        WHERE id = ?
        """,
        (
            success_rows,
            failed_rows,
            total_quantity_added,
            status,
            import_id,
        ),
    )

    conn.commit()
    conn.close()


# =========================================================
# IMPORT ERROR STORAGE
# =========================================================

def save_import_error(
    import_id,
    excel_row,
    product_id,
    product_name,
    error_message,
):
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        """
        INSERT INTO import_errors (
            import_id,
            excel_row,
            product_id,
            product_name,
            error_message,
            created_at
        )
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (
            import_id,
            excel_row,
            product_id,
            product_name,
            error_message,
            get_current_timestamp(),
        ),
    )

    conn.commit()
    conn.close()


# =========================================================
# PROCESSED EXCEL
# =========================================================

def make_processed_filename(
    original_name
):
    original_path = Path(
        original_name
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    return (
        f"{original_path.stem}"
        f"_PROCESSED_"
        f"{timestamp}.xlsx"
    )


def create_processed_excel(
    source_file,
    processed_file,
    result_rows,
):
    """
    Copy original Excel and add result columns.
    """

    shutil.copy2(
        source_file,
        processed_file,
    )

    workbook = load_workbook(
        processed_file
    )

    worksheet = workbook.active

    original_headers = [
        cell.value
        for cell in worksheet[1]
    ]

    extra_headers = [
        "Database Status",
        "Entry Number",
        "Previous Stock",
        "Quantity Added",
        "New Stock",
        "Imported At",
        "Error",
    ]

    header_map = {}

    for column_number, header in enumerate(
        original_headers,
        start=1,
    ):
        if header:
            header_map[
                str(header).strip()
            ] = column_number

    next_column = (
        worksheet.max_column + 1
    )

    extra_column_map = {}

    for header in extra_headers:

        if header in header_map:

            extra_column_map[
                header
            ] = header_map[header]

        else:

            worksheet.cell(
                row=1,
                column=next_column,
                value=header,
            )

            extra_column_map[
                header
            ] = next_column

            next_column += 1

    for result in result_rows:

        excel_row = result[
            "excel_row"
        ]

        worksheet.cell(
            row=excel_row,
            column=extra_column_map[
                "Database Status"
            ],
            value=result.get(
                "status",
                "",
            ),
        )

        worksheet.cell(
            row=excel_row,
            column=extra_column_map[
                "Entry Number"
            ],
            value=result.get(
                "entry_number",
                "",
            ),
        )

        worksheet.cell(
            row=excel_row,
            column=extra_column_map[
                "Previous Stock"
            ],
            value=result.get(
                "previous_quantity",
                "",
            ),
        )

        worksheet.cell(
            row=excel_row,
            column=extra_column_map[
                "Quantity Added"
            ],
            value=result.get(
                "quantity_added",
                "",
            ),
        )

        worksheet.cell(
            row=excel_row,
            column=extra_column_map[
                "New Stock"
            ],
            value=result.get(
                "new_quantity",
                "",
            ),
        )

        worksheet.cell(
            row=excel_row,
            column=extra_column_map[
                "Imported At"
            ],
            value=result.get(
                "imported_at",
                "",
            ),
        )

        worksheet.cell(
            row=excel_row,
            column=extra_column_map[
                "Error"
            ],
            value=result.get(
                "error",
                "",
            ),
        )

    workbook.save(
        processed_file
    )


# =========================================================
# MASTER STOCK EXCEL
# =========================================================

def update_master_stock_excel():
    """
    Create/rebuild Master_Stock.xlsx from database.

    This is used for:
    - manual stock entry
    - Excel stock import
    - latest current stock report
    """

    products = get_inventory_products()

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Current Stock"

    headers = [
        "Product ID",
        "Category",
        "Product Name",
        "Brand",
        "Model",
        "Quantity",
        "Unit",
        "Location",
        "Last Updated",
    ]

    worksheet.append(
        headers
    )

    for product in products:

        worksheet.append(
            [
                product[
                    "product_id"
                ],
                product[
                    "category"
                ],
                product[
                    "product_name"
                ],
                product[
                    "brand"
                ],
                product[
                    "model"
                ],
                product[
                    "current_quantity"
                ],
                product[
                    "unit"
                ],
                product[
                    "location"
                ],
                product[
                    "updated_at"
                ],
            ]
        )

    worksheet.freeze_panes = "A2"

    widths = {
        "A": 16,
        "B": 32,
        "C": 38,
        "D": 20,
        "E": 20,
        "F": 14,
        "G": 12,
        "H": 18,
        "I": 24,
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    master_stock_file = get_master_stock_file()

    workbook.save(
        master_stock_file
    )

    return master_stock_file


# =========================================================
# MAIN EXCEL IMPORT
# =========================================================

def import_excel_to_database(
    file_path
):
    file_path = Path(
        file_path
    )

    if not file_path.exists():
        raise ValueError(
            "Excel file does not exist."
        )

    if (
        file_path.suffix.lower()
        != ".xlsx"
    ):
        raise ValueError(
            "Please upload an .xlsx Excel file."
        )

    dataframe = read_excel_file(
        file_path
    )

    validate_excel_columns(
        dataframe
    )

    if dataframe.empty:
        raise ValueError(
            "Excel file contains no stock rows."
        )

    file_hash = generate_file_hash(
        file_path
    )

    previous_import = (
        is_file_already_imported(
            file_hash
        )
    )

    if previous_import:

        raise ValueError(
            "This exact Excel file "
            "has already been imported as "
            f"{previous_import['import_number']}."
        )

    import_id, import_number = (
        create_import_history(
            file_name=file_path.name,
            original_file_name=file_path.name,
            total_rows=len(dataframe),
            file_hash=file_hash,
        )
    )

    success_rows = 0
    failed_rows = 0
    total_quantity_added = 0.0

    result_rows = []

    for dataframe_index, row in dataframe.iterrows():

        excel_row = (
            dataframe_index + 2
        )

        product_id = clean_value(
            row.get("Product ID")
        )

        category = clean_value(
            row.get("Category")
        )

        product_name = clean_value(
            row.get("Product Name")
        )

        brand = clean_value(
            row.get("Brand")
        )

        model = clean_value(
            row.get("Model")
        )

        unit = (
            clean_value(
                row.get("Unit")
            )
            or "Nos"
        )

        location = clean_value(
            row.get("Location")
        )

        stock_date = clean_stock_date(
            row.get("Stock Date")
        )

        remarks = clean_value(
            row.get("Remarks")
        )

        try:

            if not product_id:
                raise ValueError(
                    "Product ID is required."
                )

            if not category:
                raise ValueError(
                    "Category is required."
                )

            if not product_name:
                raise ValueError(
                    "Product Name is required."
                )

            quantity = clean_quantity(
                row.get("Quantity")
            )

            result = add_stock(
                product_id=product_id,
                category=category,
                product_name=product_name,
                brand=brand,
                model=model,
                quantity_added=quantity,
                unit=unit,
                location=location,
                stock_date=stock_date,
                import_id=import_id,
                excel_row=excel_row,
                source="Excel Import",
                remarks=remarks,
            )

            success_rows += 1

            total_quantity_added += (
                quantity
            )

            result_rows.append(
                {
                    "excel_row":
                        excel_row,

                    "status":
                        "SUCCESS",

                    "entry_number":
                        result[
                            "entry_number"
                        ],

                    "previous_quantity":
                        result[
                            "previous_quantity"
                        ],

                    "quantity_added":
                        result[
                            "quantity_added"
                        ],

                    "new_quantity":
                        result[
                            "new_quantity"
                        ],

                    "imported_at":
                        get_current_timestamp(),

                    "error":
                        "",
                }
            )

        except Exception as error:

            failed_rows += 1

            error_message = str(
                error
            )

            save_import_error(
                import_id=import_id,
                excel_row=excel_row,
                product_id=product_id,
                product_name=product_name,
                error_message=(
                    error_message
                ),
            )

            result_rows.append(
                {
                    "excel_row":
                        excel_row,

                    "status":
                        "FAILED",

                    "entry_number":
                        "",

                    "previous_quantity":
                        "",

                    "quantity_added":
                        "",

                    "new_quantity":
                        "",

                    "imported_at":
                        get_current_timestamp(),

                    "error":
                        error_message,
                }
            )

    if (
        success_rows > 0
        and failed_rows == 0
    ):
        final_status = "Completed"

    elif (
        success_rows > 0
        and failed_rows > 0
    ):
        final_status = (
            "Completed With Errors"
        )

    else:
        final_status = "Failed"

    update_import_history(
        import_id=import_id,
        success_rows=success_rows,
        failed_rows=failed_rows,
        total_quantity_added=(
            total_quantity_added
        ),
        status=final_status,
    )

    processed_filename = (
        make_processed_filename(
            file_path.name
        )
    )

    processed_path = (
        PROCESSED_DIR
        / processed_filename
    )

    create_processed_excel(
        source_file=file_path,
        processed_file=processed_path,
        result_rows=result_rows,
    )

    # Keep Master_Stock.xlsx updated
    if success_rows > 0:
        update_master_stock_excel()

    return {
        "success": True,
        "import_id": import_id,
        "import_number": import_number,
        "status": final_status,
        "total_rows": len(
            dataframe
        ),
        "success_rows": (
            success_rows
        ),
        "failed_rows": (
            failed_rows
        ),
        "total_quantity_added": (
            total_quantity_added
        ),
        "processed_file": str(
            processed_path
        ),
        "master_stock_file": str(
            get_master_stock_file()
        ),
        "results": result_rows,
    }


# =========================================================
# DOWNLOADABLE EXCEL TEMPLATE
# =========================================================

def get_excel_template_dataframe():
    return pd.DataFrame(
        [
            {
                "Product ID":
                    "LAB0001",

                "Category":
                    "Laboratory Instruments",

                "Product Name":
                    "Digital pH Meter",

                "Brand":
                    "Example Brand",

                "Model":
                    "PH-100",

                "Quantity":
                    5,

                "Unit":
                    "Nos",

                "Location":
                    "Rack A1",

                "Stock Date":
                    datetime.now().strftime(
                        "%Y-%m-%d"
                    ),

                "Remarks":
                    "New stock",
            },

            {
                "Product ID":
                    "CIV0001",

                "Category":
                    "Civil Construction Instruments",

                "Product Name":
                    "Concrete Rebound Hammer",

                "Brand":
                    "Example Brand",

                "Model":
                    "RH-01",

                "Quantity":
                    2,

                "Unit":
                    "Nos",

                "Location":
                    "Rack B1",

                "Stock Date":
                    datetime.now().strftime(
                        "%Y-%m-%d"
                    ),

                "Remarks":
                    "New stock",
            },
        ]
    )


def get_excel_template_bytes():
    """
    Optional helper for Streamlit download button.
    """

    dataframe = (
        get_excel_template_dataframe()
    )

    buffer = BytesIO()

    dataframe.to_excel(
        buffer,
        index=False,
        engine="openpyxl",
    )

    buffer.seek(0)

    return buffer.getvalue()


# =========================================================
# TEST
# =========================================================

if __name__ == "__main__":

    print(
        "Excel Handler Ready"
    )

    print()
    print(
        "Required Excel Columns:"
    )

    for column in REQUIRED_COLUMNS:
        print(
            f" - {column}"
        )

    print()
    print(
        "Optional Excel Columns:"
    )

    for column in OPTIONAL_COLUMNS:
        print(
            f" - {column}"
        )

    print()
    print(
        "Upload Folder:"
    )
    print(
        UPLOAD_DIR
    )

    print()
    print(
        "Processed Folder:"
    )
    print(
        PROCESSED_DIR
    )

    print()
    print(
        "Master Stock File:"
    )
    print(
        MASTER_STOCK_FILE
    )